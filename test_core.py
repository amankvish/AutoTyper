import unittest
import os
import json
import tempfile
from core.delay_engine import DelayEngine
from core.formatter import TextFormatter
from services.settings import SettingsService
from services.file_loader import FileLoaderService, UnsupportedFileTypeError, FileLoadError
from docx import Document
import config

class TestAutoTyperCore(unittest.TestCase):
    
    def test_delay_engine_wpm_calculations(self) -> None:
        """Verify standard WPM calculations and boundaries."""
        # 60 WPM = 12 / 60 = 0.2 seconds delay per character
        delay_60 = DelayEngine.get_delay("a", 60, "normal", False, 10, 100)
        self.assertAlmostEqual(delay_60, 0.2)

        # 120 WPM = 12 / 120 = 0.1 seconds delay per character
        delay_120 = DelayEngine.get_delay("a", 120, "normal", False, 10, 100)
        self.assertAlmostEqual(delay_120, 0.1)
        
        # WPM <= 0 boundary (should fallback to WPM 60 = 0.2s)
        fallback_delay = DelayEngine.calculate_base_delay(0)
        self.assertAlmostEqual(fallback_delay, 0.2)
        fallback_delay_neg = DelayEngine.calculate_base_delay(-10)
        self.assertAlmostEqual(fallback_neg := DelayEngine.get_delay("a", -10, "normal", False, 10, 100), 0.2)

    def test_delay_engine_random_distribution(self) -> None:
        """Verify user-defined random delay bounds (in milliseconds)."""
        min_ms, max_ms = 50, 150
        for _ in range(100):
            delay = DelayEngine.get_delay("a", 60, "normal", True, min_ms, max_ms)
            # 50 ms = 0.05s, 150 ms = 0.15s
            self.assertGreaterEqual(delay, 0.05)
            self.assertLessEqual(delay, 0.15)
            
        # Ensure random delays are never negative, even if inputs are invalid
        neg_delay = DelayEngine.get_delay("a", 60, "normal", True, -50, -10)
        self.assertGreaterEqual(neg_delay, 0.0)

    def test_delay_engine_human_punctuation_pauses(self) -> None:
        """Verify Human Mode punctuation and enter pauses are added correctly."""
        base = DelayEngine.calculate_base_delay(60)  # 0.2s
        
        # Period '.' adds 0.35s delay. Base fluctuated: [0.16s, 0.24s]
        # Total delay: [0.51s, 0.59s]
        for _ in range(20):
            delay_dot = DelayEngine.get_delay(".", 60, "human", False, 10, 100)
            self.assertGreaterEqual(delay_dot, 0.51)
            self.assertLessEqual(delay_dot, 0.59)

        # Newline '\n' adds 0.50s delay (HUMAN_PAUSE_ENTER). Total: [0.66s, 0.74s]
        for _ in range(20):
            delay_nl = DelayEngine.get_delay("\n", 60, "human", False, 10, 100)
            self.assertGreaterEqual(delay_nl, 0.66)
            self.assertLessEqual(delay_nl, 0.74)

    def test_formatter_formatting_preservation(self) -> None:
        """Verify spacing, tabs, blank lines, and line endings normalization."""
        # Mix of space indentation, tabs, blank lines, carriage returns
        text = "\n\n    def test():\r\n\t\tprint('hello')\r\n\r    # Blank line above\r\n"
        
        # Normalization should change all \r\n and \r to \n, leaving spacing & tabs intact
        normalized = TextFormatter.format_for_typing(text)
        expected = "\n\n    def test():\n\t\tprint('hello')\n\n    # Blank line above\n"
        self.assertEqual(normalized, expected)
        
        # Test empty document case
        self.assertEqual(TextFormatter.format_for_typing(""), "")

    def test_settings_service_strict_validation(self) -> None:
        """Verify SettingsService strict type checking on inputs."""
        service = SettingsService()
        original_wpm = service.get("wpm")
        original_delay_enabled = service.get("random_delay_enabled")
        
        # 1. Valid settings change
        service.set("wpm", 123)
        self.assertEqual(service.get("wpm"), 123)
        
        # 2. Strict Type Safety Checks
        # wpm is int, setting True (bool) must be rejected
        service.set("wpm", True)
        self.assertEqual(service.get("wpm"), 123)
        
        # random_delay_enabled is bool, setting an int must be rejected
        service.set("random_delay_enabled", 456)
        self.assertEqual(service.get("random_delay_enabled"), original_delay_enabled)
        
        # Clean up config by restoring original values
        service.set("wpm", original_wpm)
        service.set("random_delay_enabled", original_delay_enabled)

    def test_settings_service_config_corruption_recovery(self) -> None:
        """Verify settings service automatically recovers from a corrupted config.json."""
        # 1. Write corrupted json content to config file path
        with open(config.CONFIG_FILE_PATH, 'w', encoding='utf-8') as f:
            f.write("{invalid_json_corrupted: [")
            
        try:
            # 2. Instantiate settings service. It should catch error and fallback to default dictionary
            service = SettingsService()
            self.assertEqual(service.get("wpm"), config.DEFAULT_WPM)
            self.assertEqual(service.get("theme"), "dark")
            
            # 3. Verify it saved defaults back to file
            with open(config.CONFIG_FILE_PATH, 'r', encoding='utf-8') as f:
                data = json.load(f)
                self.assertEqual(data.get("wpm"), config.DEFAULT_WPM)
        finally:
            # Clean up the config.json by restoring active defaults
            service = SettingsService()
            service.save()

    def test_file_loader_txt_md_docx(self) -> None:
        """Verify TXT, Markdown (.md), and Word (.docx) document loaders."""
        # 1. Load TXT
        with tempfile.NamedTemporaryFile(suffix=".txt", delete=False, mode="w", encoding="utf-8") as temp:
            temp.write("Plain text load test.")
            temp_path = temp.name
        try:
            content = FileLoaderService.load_file(temp_path)
            self.assertEqual(content, "Plain text load test.")
        finally:
            os.remove(temp_path)

        # 2. Load MD
        with tempfile.NamedTemporaryFile(suffix=".md", delete=False, mode="w", encoding="utf-8") as temp:
            temp.write("# Header\nMarkdown plain text.")
            temp_path = temp.name
        try:
            content = FileLoaderService.load_file(temp_path)
            self.assertEqual(content, "# Header\nMarkdown plain text.")
        finally:
            os.remove(temp_path)

        # 3. Load DOCX
        doc = Document()
        doc.add_paragraph("First paragraph.")
        doc.add_paragraph("Second paragraph.")
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as temp:
            temp_path = temp.name
        try:
            doc.save(temp_path)
            content = FileLoaderService.load_file(temp_path)
            self.assertEqual(content, "First paragraph.\nSecond paragraph.")
        finally:
            os.remove(temp_path)

    def test_file_loader_unicode_emojis(self) -> None:
        """Verify loader handles extreme Unicode character sets and emojis."""
        unicode_str = "Special Characters: 🌵 ★ ü ñ \nCyrillic: Привет \nChinese: 你好 \nEmojis: 🚀🔥🎉"
        with tempfile.NamedTemporaryFile(suffix=".txt", delete=False, mode="w", encoding="utf-8") as temp:
            temp.write(unicode_str)
            temp_path = temp.name
        try:
            content = FileLoaderService.load_file(temp_path)
            self.assertEqual(content, unicode_str)
        finally:
            os.remove(temp_path)

    def test_file_loader_unsupported_extensions(self) -> None:
        """Verify loading unsupported extensions raises UnsupportedFileTypeError."""
        with tempfile.NamedTemporaryFile(suffix=".exe", delete=False) as temp:
            temp_path = temp.name
        try:
            with self.assertRaises(UnsupportedFileTypeError):
                FileLoaderService.load_file(temp_path)
        finally:
            os.remove(temp_path)

    def test_large_document_loading_and_formatting(self) -> None:
        """Verify performance and safety on large documents (100,000+ characters)."""
        # Create 100,000 characters string
        chunk = "AutoTyper formatting preservation speed test loop. " * 2000
        self.assertGreater(len(chunk), 100000)
        
        # Test formatting speed & accuracy
        normalized = TextFormatter.format_for_typing(chunk)
        self.assertEqual(len(normalized), len(chunk))
        
        # Save and read back
        with tempfile.NamedTemporaryFile(suffix=".txt", delete=False, mode="w", encoding="utf-8") as temp:
            temp.write(chunk)
            temp_path = temp.name
        try:
            content = FileLoaderService.load_file(temp_path)
            self.assertEqual(content, chunk)
        finally:
            os.remove(temp_path)

    def test_settings_service_recent_files_capping(self) -> None:
        """Verify that recent files capped list behaves like a FIFO with max 5 files."""
        service = SettingsService()
        # Clear recent list first
        service.set("recent_files", [])
        
        paths = ["p1.txt", "p2.txt", "p3.txt", "p4.txt", "p5.txt", "p6.txt"]
        for p in paths:
            service.add_recent_file(p)
            
        recent = service.get("recent_files")
        self.assertEqual(len(recent), 5)
        # Newest first, so "p6.txt" should be first, "p1.txt" is dropped
        self.assertEqual(recent[0], os.path.normpath("p6.txt"))
        self.assertEqual(recent[-1], os.path.normpath("p2.txt"))
        self.assertNotIn(os.path.normpath("p1.txt"), recent)

    def test_file_loader_utf8_bom_and_utf16(self) -> None:
        """Verify that file loader successfully parses utf-8-sig (BOM) and utf-16 documents."""
        # 1. UTF-8-sig (with BOM)
        with tempfile.NamedTemporaryFile(suffix=".txt", delete=False, mode="w", encoding="utf-8-sig") as temp:
            temp.write("BOM test content")
            temp_path = temp.name
        try:
            content = FileLoaderService.load_file(temp_path)
            self.assertEqual(content, "BOM test content")
        finally:
            os.remove(temp_path)

        # 2. UTF-16
        with tempfile.NamedTemporaryFile(suffix=".txt", delete=False, mode="w", encoding="utf-16") as temp:
            temp.write("UTF-16 test content")
            temp_path = temp.name
        try:
            content = FileLoaderService.load_file(temp_path)
            self.assertEqual(content, "UTF-16 test content")
        finally:
            os.remove(temp_path)

    def test_remaining_time_calculation_math(self) -> None:
        """Verify estimated typing time math formulas (12 * R / W)."""
        # 100 characters remaining, WPM=60 -> 12 * 100 / 60 = 20 seconds
        rem_chars_1 = 100
        wpm_1 = 60
        sec_1 = int((12 * rem_chars_1) / wpm_1)
        self.assertEqual(sec_1, 20)

        # 250 characters remaining, WPM=120 -> 12 * 250 / 120 = 25 seconds
        rem_chars_2 = 250
        wpm_2 = 120
        sec_2 = int((12 * rem_chars_2) / wpm_2)
        self.assertEqual(sec_2, 25)

    def test_text_stats_calculations(self) -> None:
        """Verify that character, word, and line count stats are calculated correctly."""
        text = "Hello world!\nThis is a test.\n\nEnjoy typing."
        char_count = len(text)
        word_count = len(text.split()) if text.strip() else 0
        line_count = text.count('\n') + 1 if text else 0
        
        self.assertEqual(char_count, 43)
        self.assertEqual(word_count, 8)
        self.assertEqual(line_count, 4)

if __name__ == "__main__":
    unittest.main()
